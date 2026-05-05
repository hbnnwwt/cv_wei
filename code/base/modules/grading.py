import openpyxl

from modules.pipeline import LAYOUT


def _classify_question(q_num):
    """根据题号判断题型，从配置动态获取边界。"""
    _ch = LAYOUT['choice']
    _ju = LAYOUT['judge']
    choice_end = _ch['question_start'] + _ch['question_count'] - 1
    judge_end = _ju['question_start'] + _ju['question_count'] - 1
    if _ch['question_start'] <= q_num <= choice_end:
        return 'choice'
    elif _ju['question_start'] <= q_num <= judge_end:
        return 'judge'
    else:
        return 'essay'


class EssayGraderBase:
    """简答题评分器抽象基类。子类需实现 score() 方法。"""

    def score(self, question, reference, student_answer, max_score):
        """评分方法。

        Args:
            question: int, 题号
            reference: str, 参考答案
            student_answer: str, 学生答案
            max_score: int, 满分

        Returns:
            tuple: (score, max_score, feedback)
        """
        raise NotImplementedError


class DefaultEssayGrader(EssayGraderBase):
    """默认实现：简答题返回 0 分，标注需手动评分。"""

    def score(self, question, reference, student_answer, max_score):
        return 0, max_score, "需手动评分"


class GradingService:
    """评分模块：将识别结果与标准答案比对，生成评分报告。

    调用者（app.py / main.py）使用方式：
        svc = GradingService.from_xlsx('参考答案.xlsx')
        result = svc.grade(recognized_answers)
        report = svc.generate_report(result)

    其中 recognized_answers 格式：
        {
            'choice': {1: 'A', 2: 'C', ...},
            'judge':  {21: 'T', 22: 'F', ...},
            'essay':  {31: '学生写的文字'}
        }
    """

    def __init__(self, answer_key, essay_grader=None,
                 choice_score=3, judge_score=2, essay_max_score=20):
        """
        Args:
            answer_key: 标准答案字典
                {'choice': {1:'A', 2:'C', ...},
                 'judge':  {21:'T', 22:'F', ...},
                 'essay':  {31:'参考答案文本'}}
            essay_grader: 简答题评分器，默认使用 DefaultEssayGrader
            choice_score: 每道选择题分值
            judge_score: 每道判断题分值
            essay_max_score: 简答题满分
        """
        self.answer_key = answer_key
        self.essay_grader = essay_grader or DefaultEssayGrader()
        self.choice_score = choice_score
        self.judge_score = judge_score
        self.essay_max_score = essay_max_score

    @property
    def max_total(self):
        """动态计算满分。"""
        n_choice = len(self.answer_key.get('choice', {}))
        n_judge = len(self.answer_key.get('judge', {}))
        n_essay = len(self.answer_key.get('essay', {}))
        return (n_choice * self.choice_score
                + n_judge * self.judge_score
                + n_essay * self.essay_max_score)

    @classmethod
    def from_xlsx(cls, path):
        """从 参考答案.xlsx 加载标准答案。

        Args:
            path: xlsx 文件路径

        Returns:
            GradingService 实例
        """
        wb = openpyxl.load_workbook(path)
        ws = wb.active

        answer_key = {'choice': {}, 'judge': {}, 'essay': {}}
        for col in range(2, ws.max_column + 1):
            q_num = ws.cell(row=1, column=col).value
            answer = ws.cell(row=2, column=col).value
            if q_num is None or answer is None:
                continue
            try:
                q_num = int(q_num)
            except (ValueError, TypeError):
                continue
            q_type = _classify_question(q_num)
            answer_key[q_type][q_num] = answer

        _scoring = LAYOUT.get('scoring', {})
        return cls(answer_key,
                   choice_score=_scoring.get('choice_score', 3),
                   judge_score=_scoring.get('judge_score', 2),
                   essay_max_score=_scoring.get('essay_max_score', 20))

    def grade(self, recognized_answers):
        """将识别结果与标准答案比对，计算得分。

        Args:
            recognized_answers: 各题型识别结果
                {'choice': {1:'A', 2:'B', ...},
                 'judge':  {21:'T', 22:'F', ...},
                 'essay':  {31:'识别文本'}}

        Returns:
            dict: 包含各题得分、总分、错误明细的评分报告
                {
                    'choice': {题号: {'correct': 'A', 'given': 'C', 'score': 0}},
                    'judge':  {题号: {'correct': 'T', 'given': 'F', 'score': 0}},
                    'essay_detail': {题号: {'score': 0, 'max_score': 20, 'feedback': '...'}},
                    'essay': {题号: '学生文字'},
                    'choice_total': int,
                    'judge_total': int,
                    'essay_total': int,
                    'total': int,
                }

        思路提示：
            - recognized_answers 里没有某道题（key 不存在）和有但值为 None，有什么区别？
              两种情况分别应该如何给分？
            - choice_detail 和 judge_detail 的结构为什么要包含 correct、given、score 三个字段？
              如果只需要总分，given 和 correct 字段还需要保存吗？
            - essay 题的评分和 choice/judge 有什么本质区别？
              essay_grader.score() 返回的三元组 (score, max_score, feedback) 各是什么用途？
            - 学生某题没填涂（given=None），应该给 0 分还是跳过不记？
              这会影响总分计算吗？
        """
        raise NotImplementedError("TODO: 请实现评分逻辑 grade() 方法")

    def generate_report(self, result):
        """生成可读的评分报告。

        Args:
            result: grade() 的返回值

        Returns:
            str: 格式化的评分报告文本

        思路提示：
            - 如何让老师一眼看出哪道题对了、哪道题错了？
              纯数字列表和带符号（✅/❌）的列表，哪个信息密度更高？
            - 如果选择/判断题数量很多（20+10道），每道题都展开列出的报告阅读体验好吗？
              有没有什么折中的展示方式？
            - generate_report 的使用者是谁？是程序（后续处理）还是人（老师/学生）？
              这个判断会影响返回格式的选择吗？
        """
        raise NotImplementedError("TODO: 请实现报告生成 generate_report() 方法")

    def save_result_xlsx(self, template_path, output_path, student_results):
        """将识别结果按模板格式输出到 xlsx。

        Args:
            template_path: 结果.xlsx 模板路径
            output_path: 输出文件路径
            student_results: list[(student_id, recognized_answers)]
        """
        wb = openpyxl.load_workbook(template_path)
        ws = wb.active
        next_row = ws.max_row + 1

        for student_id, rec in student_results:
            ws.cell(row=next_row, column=1, value=student_id)
            choice_scores = []
            judge_scores = []
            for col in range(2, ws.max_column + 1):
                q_num = ws.cell(row=1, column=col).value
                if q_num is None:
                    continue
                try:
                    q_num = int(q_num)
                except (ValueError, TypeError):
                    continue
                q_type = _classify_question(q_num)
                if q_type == 'choice':
                    ans = rec.get('choice', {}).get(q_num)
                    ws.cell(row=next_row, column=col, value=ans)
                    correct = self.answer_key['choice'].get(q_num)
                    score = self.choice_score if ans and ans == correct else 0
                    choice_scores.append(score)
                elif q_type == 'judge':
                    ans = rec.get('judge', {}).get(q_num)
                    ws.cell(row=next_row, column=col, value=ans)
                    correct_j = self.answer_key['judge'].get(q_num)
                    score = self.judge_score if ans and ans == correct_j else 0
                    judge_scores.append(score)
            next_row += 1

            ws.cell(row=next_row, column=1, value=f"{student_id}_score")
            col_idx = 2
            for s in choice_scores:
                ws.cell(row=next_row, column=col_idx, value=s)
                col_idx += 1
            for s in judge_scores:
                ws.cell(row=next_row, column=col_idx, value=s)
                col_idx += 1
            if choice_scores:
                choice_col_start = openpyxl.utils.get_column_letter(2)
                choice_col_end = openpyxl.utils.get_column_letter(2 + len(choice_scores) - 1)
                ws.cell(row=next_row, column=33,
                        value=f"=SUM({choice_col_start}{next_row}:{choice_col_end}{next_row})")
            else:
                ws.cell(row=next_row, column=33, value=0)
            if judge_scores:
                judge_col_start = openpyxl.utils.get_column_letter(2 + len(choice_scores))
                judge_col_end = openpyxl.utils.get_column_letter(2 + len(choice_scores) + len(judge_scores) - 1)
                ws.cell(row=next_row, column=34,
                        value=f"=SUM({judge_col_start}{next_row}:{judge_col_end}{next_row})")
            else:
                ws.cell(row=next_row, column=34, value=0)
            next_row += 1

        wb.save(output_path)


if __name__ == '__main__':
    import sys
    if len(sys.argv) > 1:
        service = GradingService.from_xlsx(sys.argv[1])
        print(f"已加载标准答案: 选择题{len(service.answer_key['choice'])}题, "
              f"判断题{len(service.answer_key['judge'])}题, "
              f"简答题{len(service.answer_key['essay'])}题")
